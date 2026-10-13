from __future__ import annotations

import io
import json
import traceback
from pathlib import Path

import streamlit as st

from backend.comparison import compare_versions
from backend.documents import DefaultDocumentParser
from backend.history import HistoryRepository
from backend.i18n import Translator
from backend.models import ContentInput, RevisionInstruction, RewriteValidationResult, SuggestionDecision
from backend.providers import configured_provider
from backend.pipeline import ContentQAPipeline
from backend.policy import ReviewPolicyBuilder
from backend.revision import RevisionService, SectionRegenerator
from backend.services import ContentParser
from presentation import content_language_label, label

st.set_page_config(page_title="Content QA", page_icon=":material/auto_awesome:", layout="wide")
st.html("""
<style>
html, body, [class*="st-"], button, input, textarea, select { font-family: Inter, "PingFang SC", "Microsoft YaHei", "Noto Sans SC", -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; }
.stMainBlockContainer { max-width: 1200px; padding-top: 2rem; }
.stMarkdown, .stText, .stCaption { line-height: 1.68; }
[data-testid="stMetricValue"] { font-size: 2rem; }
[data-testid="stVerticalBlockBorderWrapper"] { border-radius: 12px; }
button { font-weight: 600; }
@media (max-width: 640px) { .stMainBlockContainer { padding: 1rem; } }
</style>
""")


def init_state():
    defaults = {"language": "zh", "result": None, "content": "", "filename": "", "decisions": {},
                "session_id": None, "versions": [], "input": {}, "result_view": "overview", "exports": False,
                "interface_mode": "product", "editorial_state": "review_pending", "workflow_view": "setup",
                "second_result": None, "second_decisions": {}, "final_result": None, "last_error": ""}
    defaults["custom_rules"] = []
    for key, value in defaults.items():
        st.session_state.setdefault(key, value)


init_state()
with st.sidebar:
    _tr = Translator(st.session_state.language)
    st.session_state.language = st.selectbox(_tr("interface.language"), ["zh", "en"],
        format_func=lambda x: "中文" if x == "zh" else "English",
        index=0 if st.session_state.language == "zh" else 1)
t = Translator(st.session_state.language)
history = HistoryRepository()


def issue_rows(result=None):
    result = result or st.session_state.result
    return [] if not result else [(area, issue) for area, diagnosis in result.diagnoses.items() for issue in diagnosis.issues]


def checks(result):
    zh = st.session_state.language == "zh"
    names = (["硬性合规", "关键词检查", "事实一致性", "人设一致性", "图文一致性", "要求遵循", "内容结构"] if zh else
             ["Hard compliance", "Keyword check", "Fact consistency", "Persona alignment", "Image-text consistency", "Instruction following", "Content structure"])
    values = [result.evaluation.compliance_pass, result.evaluation.keyword_pass, result.evaluation.fact_consistency,
              result.evaluation.persona_alignment, result.evaluation.image_text_consistency,
              result.evaluation.instruction_following, result.evaluation.structure_quality]
    return list(zip(names, values))


def summary(result):
    values = checks(result)
    return sum(ok for _, ok in values), len(values)


def accepted_ids(store="decisions"):
    return [key for key, value in st.session_state[store].items() if value.decision == "accepted"]


def editorial_state_label(state=None):
    state = state or st.session_state.editorial_state
    labels = {
        "review_pending": ("待审稿", "Pending review"), "reviewing": ("审稿中", "Reviewing"),
        "first_review_complete": ("一审完成", "First review complete"),
        "waiting_for_user_decision": ("等待确认", "Waiting for confirmation"),
        "first_revision_complete": ("一改完成", "First revision complete"),
        "second_review_complete": ("二审完成", "Second review complete"),
        "ready_to_deliver": ("可交付", "Ready to deliver"),
    }
    return labels.get(state, (state, state))[0 if st.session_state.language == "zh" else 1]


def run_review(input_data, suggestions=None):
    """Run the pipeline without leaking implementation details in Product mode."""
    try:
        st.session_state.last_error = ""
        return ContentQAPipeline("data").run(input_data, suggestions)
    except Exception:
        st.session_state.last_error = traceback.format_exc()
        return None


def show_processing_error():
    zh = st.session_state.language == "zh"
    st.error("本次处理没有完成" if zh else "This review could not be completed")
    st.write("系统在处理稿件时遇到异常，你的原稿和当前审核记录不会丢失。" if zh else
             "The system encountered a problem while processing the draft. Your source content and current review record are safe.")
    with st.container(horizontal=True):
        if st.button("重新尝试" if zh else "Try again", key="retry-processing"):
            st.rerun()
        if st.button("查看技术详情" if zh else "View technical details", key="open-error-details"):
            st.session_state.interface_mode = "developer"; st.rerun()
    if st.session_state.interface_mode == "developer" and st.session_state.last_error:
        with st.expander("查看技术详情" if zh else "View technical details"):
            st.code(st.session_state.last_error)


def area_label(area):
    return label("area", area if area in {"opening", "core", "ending", "persona"} else "multimodal", st.session_state.language, area)


def issue_copy(area, issue):
    zh = st.session_state.language == "zh"
    keyword = st.session_state.input.get("keyword", "")
    where = area_label(area)
    if issue.issue_type == "keyword":
        return ((f"当前{where}没有包含本次任务要求的关键词“{keyword}”。", "指定关键词需要自然出现在正文中，缺失会使内容未完整满足任务要求。", f"保留原有语气，自然加入“{keyword}”，避免生硬堆砌。") if zh else
                (f"The {where.lower()} does not include the required keyword “{keyword}”.", "The requested keyword must appear naturally for the copy to meet the brief.", f"Add “{keyword}” naturally without keyword stuffing."))
    if issue.issue_type == "persona":
        target = st.session_state.input.get("target_persona", "")
        return (("当前表达方式与目标人设不完全一致。", "一致的人设能让内容保持可信、连贯。", f"调整为“{target}”的语气和经验视角。") if zh else
                ("The voice does not fully match the target persona.", "A consistent persona keeps the content credible and coherent.", f"Adjust the voice for “{target}”."))
    if issue.issue_type == "product_fact":
        return (("正文与图片中的产品信息不一致。", "图文冲突会降低内容可信度。", "核对品牌、产品和规格后再生成。") if zh else
                ("Product information conflicts with the images.", "Conflicting evidence reduces trust.", "Verify the brand, product and specifications."))
    return ((f"{where}内容不完整或结构需要调整。", "完整、清晰的结构有助于读者理解重点。", "补充具体内容并保持原有表达风格。") if zh else
            (f"The {where.lower()} is incomplete or needs restructuring.", "A clear structure makes the key message easier to understand.", "Add specific detail while preserving the original voice."))


def save_version(instruction=""):
    result = st.session_state.result
    passed, total = summary(result)
    score = round(passed / total * 100, 1)
    if not st.session_state.session_id:
        st.session_state.session_id = history.create_session({"source_filename": st.session_state.filename,
            "keyword": st.session_state.input.get("keyword", ""), "content_type": st.session_state.input.get("content_type", ""),
            "detected_persona": result.parsed.persona, "generation_route": result.report["route"],
            "issue_count": len(issue_rows()), "evaluation_score": score, "provider": "mock", "model": "mock-v1"})
    history.save_version(st.session_state.session_id, {"content": result.final_text, "titles": result.titles, "tags": result.tags,
        "applied_suggestions": accepted_ids(), "user_instruction": instruction,
        "evaluation": {"score": score, "dimensions": result.evaluation.model_dump()},
        "execution_trace": [x.model_dump() for x in result.execution_trace], "input": st.session_state.input,
        "review_policy": result.review_policy.model_dump() if result.review_policy else None,
        "rewrite_validation": getattr(result, "rewrite_validation", None).model_dump() if getattr(result, "rewrite_validation", None) else None})
    st.session_state.versions = history.get_versions(st.session_state.session_id)


def render_overview(result):
    zh = st.session_state.language == "zh"
    issues = issue_rows(result)
    st.header("分析概览" if zh else "Analysis overview")
    if issues:
        st.write(f"这篇内容整体可用，但发现 **{len(issues)} 个需要修改的问题**。" if zh else f"This content is usable overall, but **{len(issues)} issues need attention**.")
        st.markdown("**主要问题**" if zh else "**Main issues**")
        for area, issue in issues[:4]:
            st.markdown(f"- {issue_copy(area, issue)[0]}")
    else:
        st.success("未发现需要修改的问题。你仍然可以继续生成表达优化版本。" if zh else "No issues require changes. You can still refine the writing.")
    if result.report.get("hard_rule_fixes"):
        st.warning("发现硬性合规表达，系统已在优化稿中处理。" if zh else "Hard-compliance wording was found and handled in the optimized draft.")
    else:
        st.success("未发现明显硬性合规问题" if zh else "No obvious hard-compliance issues found")
    (st.success if result.evaluation.structure_quality else st.warning)("核心内容结构正常" if zh and result.evaluation.structure_quality else "内容结构需要调整" if zh else "Core content structure is sound" if result.evaluation.structure_quality else "Content structure needs attention")
    cards = [(t("content.type"), label("content_direction", st.session_state.input.get("sub_direction", "日常分享"), st.session_state.language)),
             ("内容目标" if zh else "Content goal", label("content_goal", st.session_state.input.get("content_goal", "品牌认知"), st.session_state.language)),
             ("表达风格" if zh else "Writing style", label("writing_style", st.session_state.input.get("writing_style", "真实自然"), st.session_state.language)),
             ("优化程度" if zh else "Rewrite intensity", label("intensity", st.session_state.input.get("rewrite_intensity", "standard"), st.session_state.language)),
             (t("detected.persona"), result.parsed.persona),
             (t("issues.found"), f"{len(issues)} 项" if zh else str(len(issues))),
             (t("content.language"), content_language_label(st.session_state.input.get("content_language", "auto"), result.parsed.content_language, st.session_state.language))]
    if st.session_state.interface_mode == "developer": cards.append((t("generation.mode"), label("route", result.report["route"], st.session_state.language)))
    if result.vision:
        cards.append(("图片分析" if zh else "Images analyzed", f"{len(result.vision)} 张" if zh else str(len(result.vision))))
    for start in range(0, len(cards), 3):
        for column, (name, value) in zip(st.columns(min(3, len(cards) - start)), cards[start:start + 3]):
            column.metric(name, value)
    if st.button("查看并处理修改建议" if zh else "Review and handle suggestions", type="primary"):
        st.session_state.result_view = "diagnosis"; st.rerun()


def render_issue(area, issue):
    problem, reason, suggestion = issue_copy(area, issue)
    decision = st.session_state.decisions.get(issue.issue_id)
    with st.container(border=True):
        st.markdown(f"**{label('severity', issue.severity, st.session_state.language)} · {label('issue_type', issue.issue_type, st.session_state.language)}**")
        for heading, copy in [(t("problem"), problem), (t("basis"), reason), (t("evidence"), f"“{issue.evidence}”"), (t("suggested.change"), suggestion)]:
            st.markdown(f"**{heading}**"); st.write(copy)
        if issue.rule_id:
            policy_rule = next((r for r in st.session_state.result.review_policy.hard_rules + st.session_state.result.review_policy.soft_rules if r.rule_id == issue.rule_id), None) if st.session_state.result.review_policy else None
            displayed_source = label("rule_source", policy_rule.source, st.session_state.language) if policy_rule else issue.rule_source
            displayed_description = (policy_rule.description_en or policy_rule.description) if policy_rule and st.session_state.language == "en" else issue.rule_description
            st.markdown("**触发规则**" if st.session_state.language == "zh" else "**Triggered rule**"); st.write(issue.rule_id)
            st.markdown("**规则来源**" if st.session_state.language == "zh" else "**Rule source**"); st.write(displayed_source)
            st.markdown("**规则**" if st.session_state.language == "zh" else "**Rule**"); st.write(displayed_description)
        with st.container(horizontal=True):
            if st.button(t("accept"), key=f"accept-{issue.issue_id}", type="primary" if decision and decision.decision == "accepted" else "secondary"):
                st.session_state.decisions[issue.issue_id] = SuggestionDecision(issue_id=issue.issue_id, decision="accepted"); st.rerun()
            if st.button(t("ignore"), key=f"ignore-{issue.issue_id}"):
                st.session_state.decisions[issue.issue_id] = SuggestionDecision(issue_id=issue.issue_id, decision="ignored"); st.rerun()
            if st.button(t("regenerate"), key=f"regen-{issue.issue_id}"):
                st.session_state.decisions[issue.issue_id] = SuggestionDecision(issue_id=issue.issue_id, decision="regenerate"); st.rerun()
        if decision: st.caption(f"{t('status')}: {label('decision', decision.decision, st.session_state.language)}")


def render_diagnosis(result):
    zh = st.session_state.language == "zh"
    issues = issue_rows(result)
    st.header("诊断与建议" if zh else "Diagnosis & suggestions")
    if not issues: st.success("未发现需要修改的问题。" if zh else "No issues require changes.")
    for area, issue in issues: render_issue(area, issue)
    chosen = len(accepted_ids())
    st.caption(f"已接受 {chosen} 条建议。" if zh else f"{chosen} suggestions accepted.")
    if not chosen:
        st.info("请选择需要应用的修改建议。" if zh else "Choose the suggestions you want to apply."); return
    st.subheader("本次改写计划" if zh else "Rewrite plan")
    affected = {area for area, issue in issues if issue.issue_id in accepted_ids()}
    plans = [("优化开头", "自然加入指定关键词，同时保留原有分享语气。"), ("保留核心内容", "保留没有明显问题的核心表达，避免大幅结构调整。"),
             ("优化结尾", "自然回扣重点，避免明显广告语气。"), ("保持整体风格", f"保持“{result.parsed.persona}”真实、自然的分享风格。")]
    if not zh: plans = [("Improve the opening", "Add the keyword naturally while preserving the original voice."), ("Keep the core content", "Preserve sound core ideas and avoid unnecessary restructuring."), ("Improve the ending", "Reinforce the key point without an overtly promotional close."), ("Preserve the overall style", f"Keep the authentic voice of “{result.parsed.persona}”.")]
    for i, (name, copy) in enumerate(plans, 1): st.markdown(f"**{i}. {name}**"); st.write(copy)
    left, right = st.columns(2)
    changed = [area_label(x) for x in ("opening", "core", "ending", "persona") if x in affected]
    kept = [area_label(x) for x in ("opening", "core", "ending") if x not in affected]
    left.markdown(("**将修改**\n\n" if zh else "**Will change**\n\n") + "\n\n".join(f"✓ {x}" for x in changed))
    right.markdown(("**保持不变**\n\n" if zh else "**Will keep**\n\n") + "\n\n".join(f"✓ {x}" for x in kept))
    if st.button("应用已接受的建议并优化" if zh else "Apply accepted suggestions and optimize", type="primary"):
        st.session_state.result = ContentQAPipeline("data").run(st.session_state.input, accepted_ids())
        st.session_state.result_view = "result"; st.rerun()


def render_result(result):
    zh = st.session_state.language == "zh"
    st.header("优化结果" if zh else "Optimized result")
    st.subheader("修改前后对比" if zh else "Before & after")
    before, after = st.columns(2)
    with before.container(border=True, height="stretch"):
        st.markdown("**修改前**" if zh else "**Before**"); st.markdown(st.session_state.input.get("original_text", "") or "—")
    with after.container(border=True, height="stretch"):
        st.markdown("**修改后**" if zh else "**After**"); st.markdown(result.final_text or "—")
    passed, total = summary(result); validation = getattr(result, "rewrite_validation", None); validation_missing = validation is None
    if validation_missing:
        st.warning("当前版本尚未执行改写质量验证。" if zh else "Rewrite validation has not been run for this version.")
        validation = RewriteValidationResult()
    metrics = st.columns(4)
    metrics[0].metric("已应用建议" if zh else "Suggestions applied", f"{validation.accepted_suggestions_applied_count} / {validation.accepted_suggestions_total}")
    for col, (name, ok) in zip(metrics[1:], checks(result)[:3]): col.metric(name, "通过" if ok and zh else "Passed" if ok else "需关注" if zh else "Review")
    validation_cols = st.columns(2)
    validation_cols[0].metric("已解决问题" if zh else "Issues resolved", f"{validation.target_issues_resolved} / {validation.target_issues_total}")
    validation_cols[1].metric("新增问题" if zh else "New issues", len(validation.new_issues))
    if not validation_missing:
        (st.success if validation.passed else st.error)("改写质量校验通过" if zh and validation.passed else "改写未通过质量校验，请重新生成或调整建议。" if zh else "Rewrite quality validation passed" if validation.passed else "The rewrite did not pass validation. Regenerate it or adjust the selected suggestions.")
    st.markdown("**本次修改**" if zh else "**Changes in this rewrite**")
    change_lines = [issue_copy(area, issue)[2] for area, issue in issue_rows(result) if issue.issue_id in accepted_ids() or not accepted_ids()]
    for line in (change_lines[:3] or (["保留核心事实并优化整体表达。"] if zh else ["Preserved core facts and improved the overall writing."])): st.markdown(f"✓ {line}")
    st.subheader("推荐标题" if zh else "Recommended titles")
    for i, title in enumerate(result.titles, 1): st.write(f"{i}. {title}")
    st.subheader("优化后正文" if zh else "Optimized copy")
    edited = st.text_area("optimized-content", result.final_text, height=300, label_visibility="collapsed")
    if edited != result.final_text: result.final_text = edited
    st.markdown("**标签**" if zh else "**Tags**"); st.write(" ".join(result.tags) or "—")
    st.markdown("**修改说明**" if zh else "**What changed**"); st.write("已根据接受的建议优化表达，并保留原有核心信息。" if zh else "Accepted recommendations were applied while preserving the core message.")
    instruction = st.text_input(t("revision.instruction"), placeholder="让语气更自然。" if zh else "Make the tone more natural.")
    with st.container(horizontal=True):
        if st.button(t("continue")) and instruction:
            result.final_text = RevisionService().revise(result.final_text, RevisionInstruction(instruction=instruction), ContentInput.model_validate(st.session_state.input)); st.rerun()
        if st.button("重新生成" if zh else "Regenerate"):
            st.session_state.result = ContentQAPipeline("data").run(st.session_state.input, accepted_ids() or None); st.rerun()
        if st.button(t("save")):
            save_version(instruction); st.toast("版本已保存" if zh else "Version saved", icon=":material/check_circle:")
        if st.button(t("export")): st.session_state.exports = not st.session_state.exports
    with st.expander("更多调整" if zh else "More adjustments"):
        st.caption("仅重新生成需要微调的部分，其余内容保持不变。" if zh else "Regenerate only the section that needs refinement.")
        with st.container(horizontal=True):
            for section, button_text in [("opening", t("regenerate.opening")), ("core", t("regenerate.core")), ("ending", t("regenerate.ending"))]:
                if st.button(button_text, key=f"section-{section}"):
                    content_input = ContentInput.model_validate(st.session_state.input)
                    parsed = ContentParser().parse(result.final_text, content_input.target_persona, content_input.content_language)
                    regenerated = SectionRegenerator().regenerate(section, content_input, parsed, instruction)
                    result.final_text = regenerated.text
                    result.report["generation_changes"] = regenerated.changes
                    st.rerun()
            if st.button(t("regenerate.titles"), key="section-titles"):
                content_input = ContentInput.model_validate(st.session_state.input)
                parsed = ContentParser().parse(result.final_text, content_input.target_persona, content_input.content_language)
                regenerated = SectionRegenerator().regenerate_titles(content_input, parsed, instruction)
                result.titles = regenerated.titles
                result.report["generation_changes"] = regenerated.changes
                st.rerun()
    if st.session_state.exports:
        with st.container(horizontal=True):
            st.download_button(t("download.txt"), result.final_text, file_name="optimized_content.txt")
            st.download_button(t("download.markdown"), f"# {result.titles[0]}\n\n{result.final_text}\n\n{' '.join(result.tags)}", file_name="optimized_content.md")
            st.download_button(t("download.json"), json.dumps({"input": st.session_state.input, "diagnosis": result.report, "final_content": result.final_text, "evaluation": result.evaluation.model_dump()}, ensure_ascii=False, indent=2), file_name="analysis.json")
            try:
                from docx import Document
                buffer = io.BytesIO(); document = Document(); document.add_heading("Content QA Report", 0)
                document.add_heading("Optimized content", 1); document.add_paragraph(result.final_text)
                passed, total = summary(result); document.add_heading("Evaluation", 1); document.add_paragraph(f"Checks passed: {passed}/{total}")
                document.save(buffer)
                st.download_button(t("download.docx"), buffer.getvalue(), file_name="content_qa_report.docx")
            except ImportError:
                st.caption("安装 python-docx 后可导出 DOCX。" if zh else "Install python-docx to enable DOCX export.")


def render_evaluation(result):
    zh = st.session_state.language == "zh"; passed, total = summary(result)
    before_passed = max(0, total - len(issue_rows(result)) - len(result.report.get("hard_rule_fixes", [])))
    st.header("评估" if zh else "Evaluation")
    st.caption("按可解释的检查项展示，不使用无法说明来源的精确分数。" if zh else "Shown as explainable checks, not an unsupported precision score.")
    cols = st.columns(3); cols[0].metric("修改前" if zh else "Before", f"{before_passed} / {total}"); cols[1].metric("修改后" if zh else "After", f"{passed} / {total}"); cols[2].metric("提升" if zh else "Improvement", f"+{max(0, passed-before_passed)}")
    for name, ok in checks(result): (st.success if ok else st.warning)(("通过 · " if ok and zh else "Passed · " if ok else "需关注 · " if zh else "Review · ") + name)


def render_technical(result):
    zh = st.session_state.language == "zh"; st.header("技术详情" if zh else "Technical details")
    provider = result.report.get("provider", {"name": "Unknown", "model": "", "mode": "unknown"})
    if result.vision:
        with st.expander("图片理解" if zh else "Vision analysis"):
            for image in result.vision: st.write(f"{image.image_id}: {image.scene} · {image.style}")
    with st.expander("执行轨迹" if zh else "Execution trace"):
        rows = [{("步骤" if zh else "Step"): label("trace_step", x.step, st.session_state.language), ("状态" if zh else "Status"): label("status", x.status, st.session_state.language), ("耗时" if zh else "Duration"): f"{x.latency_ms:.0f} ms", ("结果" if zh else "Result"): x.error or ("完成" if zh else "Completed")} for x in result.execution_trace]
        st.dataframe(rows, hide_index=True, width="stretch")
    with st.expander("Provider 信息" if zh else "Provider information"):
        st.write(f"Provider: {provider['name']}"); st.write(f"Model: {provider['model']}"); st.write(f"Mode: {'Demo' if provider['mode'] == 'demo' else 'Live'}"); st.write(f"Token usage: {result.evaluation.token_usage}")
    with st.expander("原始结构化输出" if zh else "Raw structured output"): st.json(result.report)
    with st.expander("ReviewPolicy JSON"):
        st.json(result.review_policy.model_dump() if result.review_policy else {})
    with st.expander("Developer Details"):
        st.write(f"Route: {result.report['route']}"); st.write(result.report.get("generation_changes", []))
    render_benchmark_lab()


def render_benchmark_lab():
    zh = st.session_state.language == "zh"; st.subheader("Benchmark 与评估实验室" if zh else "Benchmark & evaluation lab")
    rows = [json.loads(line) for line in Path("benchmark/dataset.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
    counts = {}
    for row in rows: counts[row.get("category", "other")] = counts.get(row.get("category", "other"), 0) + 1
    st.metric("Benchmark Dataset", f"{len(rows)} cases")
    st.dataframe([{"Category": key.replace("_", " ").title(), "Cases": value} for key, value in sorted(counts.items())], hide_index=True, width="stretch")
    with st.container(horizontal=True):
        if st.button("查看测试集" if zh else "View dataset", key="benchmark-dataset"): st.dataframe(rows, width="stretch")
        if st.button("运行 Benchmark" if zh else "Run benchmark", key="benchmark-run"):
            from benchmark.runner import run
            with st.spinner("正在运行…" if zh else "Running…"): st.session_state.benchmark_result = run()
        if st.button("查看历史结果" if zh else "View past results", key="benchmark-history"): st.info("尚未实现 Benchmark 历史结果存储。" if zh else "Benchmark result history is not implemented yet.")
    if st.session_state.get("benchmark_result"): st.json(st.session_state.benchmark_result)


def render_workspace():
    result = st.session_state.result; zh = st.session_state.language == "zh"
    if not result:
        st.info("上传或粘贴内容后，系统将检查关键词、结构、合规、人设，以及图文一致性（如有图片）。" if zh else "Upload or paste content to check keywords, structure, compliance, persona, and image-text consistency."); return
    if result.report.get("provider", {}).get("mode") == "demo": st.caption("Demo Mode" if not zh else "演示模式")
    names = {"overview": "概览" if zh else "Overview", "diagnosis": "诊断与建议" if zh else "Diagnosis & suggestions", "result": "优化结果" if zh else "Optimized result", "evaluation": "评估" if zh else "Evaluation"}
    if st.session_state.interface_mode == "developer": names["technical"] = "技术详情" if zh else "Technical details"
    chosen = st.segmented_control("结果视图" if zh else "Result view", list(names.values()), default=names.get(st.session_state.result_view, names["overview"]), label_visibility="collapsed")
    view = next(key for key, value in names.items() if value == chosen); st.session_state.result_view = view
    {"overview": render_overview, "diagnosis": render_diagnosis, "result": render_result, "evaluation": render_evaluation, "technical": render_technical}[view](result)


def render_review_basis(content_data):
    zh = st.session_state.language == "zh"
    st.subheader("审核依据" if zh else "Review basis")
    with st.expander("+ 添加审核规则" if zh else "+ Add review rule"):
        cols = st.columns(2)
        name = cols[0].text_input("规则名称" if zh else "Rule name", key="custom-name")
        description = cols[1].text_input("规则描述" if zh else "Rule description", key="custom-description")
        rule_type = cols[0].selectbox("规则类型" if zh else "Rule type", ["硬性规则", "软性规则"] if zh else ["Hard rule", "Soft rule"], key="custom-type")
        severity = cols[1].selectbox("严重程度" if zh else "Severity", ["低", "中", "高"] if zh else ["Low", "Medium", "High"], index=1, key="custom-severity")
        scope = st.selectbox("适用范围" if zh else "Scope", ["全文", "标题", "开头", "核心内容", "结尾", "标签"] if zh else ["Full text", "Title", "Opening", "Core content", "Ending", "Tags"], key="custom-scope")
        if st.button("添加规则" if zh else "Add rule", key="add-custom-rule") and name and description:
            number = len(st.session_state.custom_rules) + 1
            type_value = "hard" if rule_type in {"硬性规则", "Hard rule"} else "soft"
            severity_value = {"低": "low", "中": "medium", "高": "high", "Low": "low", "Medium": "medium", "High": "high"}[severity]
            scope_value = {"全文": "full_text", "标题": "title", "开头": "opening", "核心内容": "core", "结尾": "ending", "标签": "tags", "Full text": "full_text", "Title": "title", "Opening": "opening", "Core content": "core", "Ending": "ending", "Tags": "tags"}[scope]
            st.session_state.custom_rules.append({"rule_id": f"CR-{number:03d}", "name": name, "description": description, "type": type_value, "severity": severity_value, "scope": scope_value, "enabled": True})
            st.rerun()
    for index, rule in enumerate(list(st.session_state.custom_rules)):
        with st.container(horizontal=True, vertical_alignment="center"):
            rule["enabled"] = st.checkbox(rule["name"], value=rule.get("enabled", True), key=f"custom-enabled-{rule['rule_id']}")
            if st.button("移除" if zh else "Remove", key=f"custom-remove-{rule['rule_id']}"):
                st.session_state.custom_rules.pop(index); st.rerun()
    content_data = {**content_data, "custom_rules": st.session_state.custom_rules}
    policy = ReviewPolicyBuilder("data").build(ContentInput.model_validate(content_data))
    source_names = {"task": "本次任务要求" if zh else "Task requirements", "brand": "品牌规则" if zh else "Brand rules", "product": "产品知识" if zh else "Product knowledge", "platform": "平台示例规则" if zh else "Demo platform rules", "rubric": "内容质量标准" if zh else "Content quality rubric", "custom": "自定义规则" if zh else "Custom rules"}
    all_rules = policy.hard_rules + policy.soft_rules
    for source in policy.sources:
        source_rules = [rule for rule in all_rules if rule.source == source["source"]]
        mark = "✓" if source_rules else "○"
        with st.expander(f"{mark} {source_names[source['source']]} · v{source['version']} · {len(source_rules)} " + ("条" if zh else "rules")):
            if source.get("source_type") == "demo": st.caption("示例规则，并非官方平台规则。" if zh else "Demo rules; not official platform policy.")
            for rule in source_rules:
                display_name = (rule.name_en or rule.name) if not zh else rule.name
                display_description = (rule.description_en or rule.description) if not zh else rule.description
                title = f"{rule.rule_id} · {display_name}" if st.session_state.interface_mode == "developer" else display_name
                st.markdown(f"**{title}**")
                st.caption(("硬性规则" if rule.type == "hard" else "软性规则") if zh else ("Hard rule" if rule.type == "hard" else "Soft rule"))
                st.write(display_description)
    if policy.conflicts:
        st.error((f"发现 {len(policy.conflicts)} 个规则冲突，需要人工确认。" if zh else f"{len(policy.conflicts)} rule conflicts require review."))
    st.caption((f"规则版本：{policy.policy_id}。诊断、改写与最终评估将使用同一快照。" if zh else f"Rule version: {policy.policy_id}. Diagnosis, rewrite, and final evaluation use this same snapshot."))
    return policy


def render_new_analysis():
    zh = st.session_state.language == "zh"
    st.title("AI 内容审稿与优化助手" if zh else "AI Content Review & Optimization Copilot")
    st.caption("帮你按照明确的审核规则，完成一审、改稿、二审和最终交付。" if zh else
               "Review, revise, re-check and finalize content using explicit review criteria.")
    with st.container(border=True):
        type_zh = ["日常分享", "产品体验", "知识科普", "场景种草", "产品测评", "对比推荐", "问题解决", "新品介绍", "活动内容", "FAQ / 答疑"]
        type_en = ["Daily sharing", "Product experience", "Educational content", "Scenario recommendation", "Product review", "Comparison", "Problem solving", "Product launch", "Campaign content", "FAQ"]
        goals = ["品牌认知", "产品介绍", "知识教育", "种草", "转化", "用户互动"] if zh else ["Brand awareness", "Product introduction", "Education", "Recommendation", "Conversion", "Engagement"]
        styles = ["真实自然", "专业可信", "轻松活泼", "简洁直接", "温暖生活化"] if zh else ["Authentic", "Professional", "Lighthearted", "Concise", "Warm lifestyle"]
        intensities = ["轻度润色", "正常优化", "深度改写"] if zh else ["Light", "Standard", "Deep"]
        config = st.columns(3)
        shown_type = config[0].selectbox("内容类型" if zh else "Content type", type_zh if zh else type_en)
        shown_goal = config[1].selectbox("内容目标" if zh else "Content goal", goals)
        shown_style = config[2].selectbox("表达风格" if zh else "Writing style", styles)
        type_index = (type_zh if zh else type_en).index(shown_type); canonical_type = type_zh[type_index]
        goal_zh = ["品牌认知", "产品介绍", "知识教育", "种草", "转化", "用户互动"]
        style_zh = ["真实自然", "专业可信", "轻松活泼", "简洁直接", "温暖生活化"]
        goal = goal_zh[goals.index(shown_goal)]; style = style_zh[styles.index(shown_style)]
        direction = "纯路人" if canonical_type in {"产品测评", "对比推荐", "FAQ / 答疑"} else "生活方式"
        intensity_shown = st.segmented_control("优化程度" if zh else "Rewrite intensity", intensities, default=intensities[1])
        intensity = ["light", "standard", "deep"][intensities.index(intensity_shown)]
        details = st.columns(4); brand = details[0].text_input(t("brand"), "示例品牌"); product = details[1].text_input(t("product"), "示例产品")
        persona = details[2].text_input(t("persona"), "新手妈妈" if zh else "First-time mother"); keyword = details[3].text_input(t("keyword"), "轻盈奶粉" if zh else "light formula")
        lang_options = [t("content.language.auto"), t("content.language.zh"), t("content.language.en")]; content_lang = st.selectbox(t("language.content"), lang_options)
        source_mode = st.segmented_control(t("source"), [t("paste"), t("upload")], default=t("paste")); filename = ""
        if source_mode == t("upload"):
            upload = st.file_uploader(t("document"), type=["txt", "md", "docx", "pdf", "csv", "xlsx"])
            if upload:
                parsed = DefaultDocumentParser().parse(upload.name, upload.getvalue()); filename = parsed.filename
                (st.success if parsed.status == "parsed" else st.error)(f"{filename} · " + ("解析完成" if zh and parsed.status == "parsed" else "解析失败" if zh else "Parsed" if parsed.status == "parsed" else "Failed")); source = st.text_area(t("content"), parsed.text, height=260)
            else: source = st.text_area(t("content"), st.session_state.content, height=260)
        else:
            example = "#早餐\n今天分享一个日常场景。\n轻盈奶粉体验不错。\n推荐按需选择。" if zh else "#Breakfast\nHere is a simple everyday moment.\nThe light formula worked well.\nChoose what fits your needs."
            source = st.text_area(t("content"), st.session_state.content or example, height=260)
        images = st.file_uploader(t("images"), type=["png", "jpg", "jpeg", "webp"], accept_multiple_files=True)
        if len(images) > 15: st.warning("仅分析前 15 张图片。" if zh else "Only the first 15 images will be analyzed.")
        code = {t("content.language.auto"): "auto", t("content.language.zh"): "zh", t("content.language.en"): "en"}[content_lang]
        input_data = {"brand": brand, "product": product, "content_type": direction, "sub_direction": canonical_type,
            "content_goal": goal, "writing_style": style, "rewrite_intensity": intensity, "target_persona": persona,
            "keyword": keyword, "original_text": source, "images": [x.name for x in images[:15]], "content_language": code,
            "custom_rules": st.session_state.custom_rules}
        policy = render_review_basis(input_data)
        st.subheader("已识别本次审稿任务" if zh else "Review task understood")
        with st.container(border=True):
            task_rows = [
                ("稿件" if zh else "Draft", filename or ("粘贴文本" if zh else "Pasted text")),
                ("图片" if zh else "Images", f"{min(len(images), 15)} 张" if zh else str(min(len(images), 15))),
                ("内容类型" if zh else "Content type", shown_type),
                ("内容目标" if zh else "Content goal", shown_goal),
                ("关键词" if zh else "Keyword", keyword or "—"),
                ("表达风格" if zh else "Writing style", shown_style),
            ]
            for name, value in task_rows: st.markdown(f"**{name}**  {value}")
            st.caption("预计审核：关键词要求、品牌规则、产品事实、合规表达、内容结构、人设风格和图文一致性。" if zh else
                       "Planned checks: keyword requirements, brand rules, product facts, compliance, structure, persona, and image-text consistency.")
        run = st.button("确认并开始审稿" if zh else "Confirm and start review", type="primary", disabled=not source.strip())
    if run:
        st.session_state.content = source; st.session_state.filename = filename
        st.session_state.input = {**input_data, "review_policy": policy.model_dump()}
        st.session_state.editorial_state = "reviewing"
        with st.spinner("正在审稿…" if zh else "Reviewing…"):
            result = run_review(st.session_state.input)
        if result:
            st.session_state.result = result; st.session_state.decisions = {}; st.session_state.second_result = None
            st.session_state.second_decisions = {}; st.session_state.final_result = None; st.session_state.session_id = None
            st.session_state.versions = []; st.session_state.editorial_state = "waiting_for_user_decision" if issue_rows(result) else "first_review_complete"
            st.session_state.workflow_view = "first"; st.rerun()
        show_processing_error()
    render_workspace()


def render_editor_issue(area, issue, number, result, store="decisions", prefix="first"):
    zh = st.session_state.language == "zh"
    problem, reason, suggestion = issue_copy(area, issue)
    decision = st.session_state[store].get(issue.issue_id)
    with st.container(border=True):
        st.subheader((f"问题 {number} · " if zh else f"Issue {number} · ") + label("issue_type", issue.issue_type, st.session_state.language))
        for heading, copy in [
            ("问题是什么" if zh else "What needs attention", problem),
            ("为什么需要改" if zh else "Why it matters", reason),
            ("原文" if zh else "Original wording", f'“{issue.evidence}”'),
            ("建议修改" if zh else "Suggested revision", suggestion),
        ]:
            st.markdown(f"**{heading}**"); st.write(copy)
        if issue.rule_id:
            policy = result.review_policy
            rule = next((r for r in policy.hard_rules + policy.soft_rules if r.rule_id == issue.rule_id), None) if policy else None
            rule_name = ((rule.name_en or rule.name) if not zh else rule.name) if rule else label("issue_type", issue.issue_type, st.session_state.language)
            st.markdown("**审核依据**" if zh else "**Review criterion**"); st.write(f"{rule_name} · {issue.rule_id}")
        with st.container(horizontal=True):
            if st.button("采用这个修改" if zh else "Use this revision", key=f"{prefix}-accept-{issue.issue_id}", type="primary" if decision and decision.decision == "accepted" else "secondary"):
                st.session_state[store][issue.issue_id] = SuggestionDecision(issue_id=issue.issue_id, decision="accepted"); st.rerun()
            if st.button("忽略" if zh else "Ignore", key=f"{prefix}-ignore-{issue.issue_id}"):
                st.session_state[store][issue.issue_id] = SuggestionDecision(issue_id=issue.issue_id, decision="ignored"); st.rerun()
            if st.button("换一个说法" if zh else "Try another wording", key=f"{prefix}-regen-{issue.issue_id}"):
                st.session_state[store][issue.issue_id] = SuggestionDecision(issue_id=issue.issue_id, decision="regenerate"); st.rerun()


def render_draft(result, final=False):
    zh = st.session_state.language == "zh"
    st.subheader("最终标题" if final and zh else "Final titles" if final else "推荐标题" if zh else "Recommended titles")
    for index, title in enumerate(result.titles, 1): st.write(f"{index}. {title}")
    st.subheader("最终正文" if final and zh else "Final content" if final else "优化后正文" if zh else "Optimized content")
    edited = st.text_area("正文" if zh else "Content", result.final_text, height=320, key="final-content" if final else "first-optimized-content", label_visibility="collapsed")
    if edited != result.final_text: result.final_text = edited
    st.subheader("最终标签" if final and zh else "Final tags" if final else "标签" if zh else "Tags")
    st.write(" ".join(result.tags) or "—")


def render_first_review(result):
    zh = st.session_state.language == "zh"; issues = issue_rows(result); automatic = len(result.report.get("hard_rule_fixes", []))
    st.header("一审完成" if zh else "First review complete")
    st.write("我已经根据本次审核要求完成第一轮审稿。" if zh else "The first editorial review is complete using the agreed review criteria.")
    cols = st.columns(3); cols[0].metric("已自动处理" if zh else "Handled automatically", f"{automatic} 项" if zh else automatic); cols[1].metric("建议确认" if zh else "Needs confirmation", f"{len(issues)} 项" if zh else len(issues)); cols[2].metric("当前状态" if zh else "Current status", editorial_state_label())
    with st.expander("查看已自动处理的内容" if zh else "See automatic fixes", expanded=bool(automatic)):
        for item in result.report.get("hard_rule_fixes", []): st.write(f"✓ {item}")
        if not automatic: st.caption("没有需要自动处理的确定性问题。" if zh else "No deterministic fixes were needed.")
    st.subheader("优化初稿" if zh else "Optimized first draft"); st.caption("先阅读完整稿件，再逐项确认编辑建议。" if zh else "Read the draft first, then decide on each suggestion."); render_draft(result)
    st.subheader("建议确认" if zh else "Suggestions to confirm")
    for number, (area, issue) in enumerate(issues, 1): render_editor_issue(area, issue, number, result)
    if not issues: st.success("没有需要确认的问题。" if zh else "Nothing needs confirmation.")
    if st.button("应用已接受建议并生成优化稿" if zh else "Apply accepted suggestions", type="primary"):
        with st.spinner("正在生成优化稿…" if zh else "Creating the optimized draft…"): revised = run_review(st.session_state.input, accepted_ids() or None)
        if revised:
            st.session_state.result = revised; st.session_state.editorial_state = "first_revision_complete"; st.session_state.workflow_view = "optimized"; st.rerun()
        show_processing_error()


def render_optimized_draft(result):
    zh = st.session_state.language == "zh"; st.header("优化稿件" if zh else "Optimized draft"); render_draft(result)
    st.subheader("本次主要调整" if zh else "What changed")
    for item in result.report.get("generation_changes", [])[:6] or (["保留核心事实，并优化整体表达。"] if zh else ["Preserved core facts and improved the overall writing."]): st.write(f"✓ {item}")
    validation = getattr(result, "rewrite_validation", None)
    if validation: (st.success if validation.passed else st.warning)("改写质量检查已通过" if zh and validation.passed else "改写质量检查未完全通过，请在二审中继续确认。" if zh else "Rewrite quality checks passed" if validation.passed else "Some quality checks need attention in the second review.")
    if st.button("开始二审" if zh else "Start second review", type="primary"):
        with st.spinner("正在进行二审…" if zh else "Running the second review…"): second = run_review({**st.session_state.input, "original_text": result.final_text})
        if second:
            st.session_state.second_result = second; st.session_state.second_decisions = {}; st.session_state.editorial_state = "second_review_complete"; st.session_state.workflow_view = "second"; st.rerun()
        show_processing_error()


def render_second_review():
    zh = st.session_state.language == "zh"; result = st.session_state.second_result
    if not result: st.info("请先完成优化稿件，再开始二审。" if zh else "Complete the optimized draft before starting the second review."); return
    issues = issue_rows(result); st.header("二审建议" if zh else "Second review"); st.write(f"还有 {len(issues)} 个地方建议确认。" if zh else f"{len(issues)} items remain for confirmation.")
    for number, (area, issue) in enumerate(issues, 1): render_editor_issue(area, issue, number, result, "second_decisions", "second")
    if not issues: st.success("二审没有发现新的待处理问题。" if zh else "The second review found no new issues.")
    with st.container(horizontal=True):
        if st.button("应用已接受建议并二改" if zh else "Apply accepted suggestions", type="primary"):
            final = run_review({**st.session_state.input, "original_text": st.session_state.result.final_text}, accepted_ids("second_decisions") or None)
            if final: st.session_state.final_result = final; st.session_state.editorial_state = "ready_to_deliver"; st.session_state.workflow_view = "final"; st.rerun()
            show_processing_error()
        if st.button("当前版本已满意" if zh else "Keep current version"):
            st.session_state.final_result = st.session_state.result; st.session_state.editorial_state = "ready_to_deliver"; st.session_state.workflow_view = "final"; st.rerun()


def render_final_version():
    zh = st.session_state.language == "zh"; result = st.session_state.final_result or st.session_state.result
    st.header("最终版本" if zh else "Final version"); st.success("状态：可交付" if zh else "Status: Ready to deliver", icon=":material/check_circle:"); render_draft(result, final=True)
    st.subheader("本次修改摘要" if zh else "Revision summary")
    for item in result.report.get("generation_changes", [])[:6] or (["已完成一审、改稿和二审。"] if zh else ["First review, revision, and second review completed."]): st.write(f"✓ {item}")
    st.subheader("最终审核结果" if zh else "Final review result")
    for name, ok in checks(result): (st.success if ok else st.warning)(("通过 · " if ok and zh else "Passed · " if ok else "需关注 · " if zh else "Review · ") + name)
    with st.container(horizontal=True):
        st.download_button("导出 Markdown" if zh else "Export Markdown", f"# {result.titles[0]}\n\n{result.final_text}\n\n{' '.join(result.tags)}", file_name="final_content.md")
        try:
            from docx import Document
            buffer = io.BytesIO(); doc = Document(); doc.add_heading(result.titles[0], 0); doc.add_paragraph(result.final_text); doc.save(buffer)
            st.download_button("导出 Word" if zh else "Export Word", buffer.getvalue(), file_name="final_content.docx")
        except ImportError: pass
        if st.button("保存版本" if zh else "Save version"):
            previous = st.session_state.result; st.session_state.result = result; save_version("最终可交付版本" if zh else "Final deliverable"); st.session_state.result = previous; st.toast("版本已保存" if zh else "Version saved", icon=":material/check_circle:")
        if st.button("继续修改" if zh else "Continue editing"): st.session_state.workflow_view = "optimized"; st.rerun()


def render_developer_workspace(result):
    views = ["Evaluation", "Benchmark", "Execution Trace", "Review Policy", "Provider", "Raw Results"]
    view = st.segmented_control("Developer view", views, default="Evaluation", label_visibility="collapsed")
    if view == "Evaluation": render_evaluation(result); st.subheader("RewriteValidationResult"); st.json(result.rewrite_validation.model_dump() if result.rewrite_validation else {})
    elif view == "Benchmark": render_benchmark_lab()
    elif view == "Execution Trace": st.dataframe([x.model_dump() for x in result.execution_trace], hide_index=True, width="stretch")
    elif view == "Review Policy": st.json(result.review_policy.model_dump() if result.review_policy else {})
    elif view == "Provider": st.json({**result.report.get("provider", {}), "token_usage": result.evaluation.token_usage})
    else: st.json({"report": result.report, "pipeline_result": result.model_dump()})


def render_workspace():
    result = st.session_state.result; zh = st.session_state.language == "zh"
    if not result: st.info("上传或粘贴内容后即可开始审稿。" if zh else "Upload or paste content to start a review."); return
    if getattr(result, "rewrite_validation", None) is None:
        st.warning("当前版本尚未执行改写质量验证。" if zh else "Rewrite validation has not been run for this version.")
    if st.session_state.interface_mode == "developer": render_developer_workspace(result); return
    steps = {"setup": "稿件信息" if zh else "Review setup", "first": "一审结果" if zh else "First review", "optimized": "优化稿件" if zh else "Optimized draft", "second": "二审建议" if zh else "Second review", "final": "最终版本" if zh else "Final version"}
    selected = st.segmented_control("审稿流程" if zh else "Review workflow", list(steps.values()), default=steps.get(st.session_state.workflow_view, steps["first"]), label_visibility="collapsed")
    view = next(key for key, value in steps.items() if value == selected); available = {"setup", "first"}
    if st.session_state.editorial_state in {"first_revision_complete", "second_review_complete", "ready_to_deliver"}: available.add("optimized")
    if st.session_state.second_result: available.add("second")
    if st.session_state.final_result: available.add("final")
    if view not in available: st.info("请先完成前一步。" if zh else "Complete the previous step first."); return
    st.session_state.workflow_view = view
    if view == "setup": st.header("稿件信息" if zh else "Review setup"); st.write(st.session_state.filename or ("粘贴文本" if zh else "Pasted text")); st.caption(f"{editorial_state_label()} · {st.session_state.input.get('sub_direction', '')} · {st.session_state.input.get('writing_style', '')}")
    elif view == "first": render_first_review(result)
    elif view == "optimized": render_optimized_draft(result)
    elif view == "second": render_second_review()
    else: render_final_version()


def render_history():
    st.title(t("nav.history")); sessions = history.list_sessions()
    if not sessions: st.info(t("history.empty")); return
    selected = st.selectbox(t("session"), sessions, format_func=lambda x: f"{x.keyword or 'Untitled'} · {x.updated_time[:19]}"); versions = history.get_versions(selected.session_id)
    if versions:
        a, b = st.columns(2); va = a.selectbox("版本 A" if st.session_state.language == "zh" else "Version A", versions, format_func=lambda x: f"{t('version')} {x.version_number}"); vb = b.selectbox("版本 B" if st.session_state.language == "zh" else "Version B", versions, index=len(versions)-1, format_func=lambda x: f"{t('version')} {x.version_number}")
        if st.button(t("compare")): st.code(compare_versions(va, vb)["diff"] or ("没有内容差异" if st.session_state.language == "zh" else "No content differences"))
        reopen = st.selectbox("重新打开版本" if st.session_state.language == "zh" else "Reopen version", versions, format_func=lambda x: f"{t('version')} {x.version_number}", key="reopen-version")
        if st.button("打开所选版本" if st.session_state.language == "zh" else "Reopen selected version"):
            st.session_state.content = reopen.content; st.session_state.filename = selected.source_filename
            st.session_state.input = {**reopen.input, "original_text": reopen.content}
            st.session_state.result = None; st.session_state.session_id = selected.session_id; st.rerun()
        st.subheader("版本" if st.session_state.language == "zh" else "Versions")
        for version in versions:
            with st.container(border=True):
                applied = len(version.applied_suggestions)
                st.write(f"{t('version')} {version.version_number} · {version.created_time[:19]} · " + (f"已应用 {applied} 条建议" if st.session_state.language == "zh" else f"{applied} suggestions applied"))
                st.caption(version.user_instruction or ("原始分析" if st.session_state.language == "zh" else "Original analysis"))
                if version.review_policy:
                    custom_count = len([r for r in version.review_policy.hard_rules + version.review_policy.soft_rules if r.source == "custom"])
                    st.caption((f"审核规则版本：{version.review_policy.policy_id} · 自定义规则：{custom_count}" if st.session_state.language == "zh" else f"Review policy: {version.review_policy.policy_id} · Custom rules: {custom_count}"))
                    with st.expander("查看规则快照" if st.session_state.language == "zh" else "Inspect policy snapshot"):
                        for source in version.review_policy.sources: st.write(f"{source['name']} v{source['version']}")
                        if st.session_state.interface_mode == "developer": st.json(version.review_policy.model_dump())


def render_settings():
    zh = st.session_state.language == "zh"; st.title(t("nav.settings"))
    labels = ["普通模式", "技术模式"] if zh else ["Product mode", "Developer mode"]
    current = 0 if st.session_state.interface_mode == "product" else 1
    selected = st.segmented_control("界面模式" if zh else "Interface mode", labels, default=labels[current])
    st.session_state.interface_mode = "product" if selected == labels[0] else "developer"
    st.caption("普通模式隐藏执行轨迹、原始数据和 Benchmark。" if zh else "Product mode hides execution traces, raw data, and Benchmark tools.")
    if st.session_state.interface_mode == "developer":
        st.info("技术模式已开启。分析完成后可在“技术详情”中查看工程信息与 Benchmark。" if zh else "Developer mode is on. Technical details and Benchmark tools appear after analysis.")
        provider = configured_provider(); st.write(f"{t('settings.provider')}: {provider.provider_name}"); st.write(f"{t('settings.model')}: {provider.model}"); st.write(f"Mode: {'Demo' if provider.mode == 'demo' else 'Live'}")


with st.sidebar:
    nav_items = ["新建审稿", "历史记录", "设置"] if st.session_state.language == "zh" else ["New Review", "History", "Settings"]
    page = st.selectbox("导航" if st.session_state.language == "zh" else "Navigation", nav_items)
if page == nav_items[0]: render_new_analysis()
elif page == nav_items[1]: render_history()
else: render_settings()
