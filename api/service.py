from __future__ import annotations

import io
import re
import uuid
from datetime import datetime, timezone

from backend.models import ContentInput, PipelineResult, ReviewPolicy
from backend.pipeline import ContentQAPipeline
from backend.policy import ReviewPolicyBuilder

from .models import ReviewCreateRequest, ReviewIssueViewModel, ReviewSummaryViewModel, WorkflowState
from .store import ReviewStore


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


ISSUE_LABELS = {
    "compliance": "合规性", "keyword": "关键词要求", "persona": "人设与语气",
    "structure": "内容结构", "product_fact": "产品事实", "image_text": "图文一致性",
}


def issue_view(issue, policy: ReviewPolicy, decisions: dict[str, str]) -> ReviewIssueViewModel:
    rule = next((r for r in policy.hard_rules + policy.soft_rules if r.rule_id == issue.rule_id), None)
    mandatory = bool(rule and rule.type == "hard") or issue.severity == "high"
    return ReviewIssueViewModel(
        id=issue.issue_id,
        priorityLabel="必须修改" if mandatory else "建议优化",
        categoryLabel=ISSUE_LABELS.get(issue.issue_type, "内容质量"),
        title=(rule.name if rule else issue.explanation) or "内容需要调整",
        description=issue.explanation or issue.rule_description,
        originalText=issue.evidence,
        suggestedText=issue.suggestion,
        ruleReference=f"{rule.source_name} · {rule.rule_id}" if rule else issue.rule_id,
        decision=decisions.get(issue.issue_id),
        isMandatory=mandatory,
    )


def unsupported_claims(source: str, candidate: str) -> list[str]:
    """Detect new high-risk factual claims that have no source support."""
    patterns = [
        r"\d+(?:\.\d+)?%", r"N\s*=\s*[\d,]+", r"\d{4}年", r"连续[一二三四五六七八九十\d]+年",
        r"(?:获得|荣获|蝉联).{0,18}(?:奖|认证)", r"超过[\d,.]+(?:万|亿)?(?:名|位|个)?用户",
    ]
    found: list[str] = []
    for pattern in patterns:
        for claim in re.findall(pattern, candidate):
            if claim not in source:
                found.append(claim)
    return sorted(set(found))


def evidence_issues(text: str) -> list[dict]:
    """Create mandatory review items for high-risk claims that need evidence."""
    issues: list[dict] = []
    if re.search(r"保证.{0,8}收益|收益率.{0,8}(?:不低于|至少)", text):
        issues.append(ReviewIssueViewModel(id="evidence_guaranteed_return", priorityLabel="必须修改", categoryLabel="合规性", title="使用了收益承诺表达", description="金融内容不能承诺收益或暗示投资结果确定。", originalText=text, suggestedText="改为客观描述历史表现，并补充“投资有风险，请审慎决策”。", ruleReference="合规规范 · 收益承诺", isMandatory=True).model_dump())
    for index, statistic in enumerate(re.findall(r"(?:超过)?\d+(?:\.\d+)?%[^。！!?]*", text)):
        issues.append(ReviewIssueViewModel(id=f"evidence_stat_{index}", priorityLabel="必须修改", categoryLabel="数据引用", title="统计数据缺少可验证来源", description="百分比、调查结论等数据必须有已批准证据，不能仅用于增强可信度。", originalText=statistic, suggestedText="删除无来源统计数据，或补充经核验的数据来源。", ruleReference="内容准确性规范 · 数据证据", isMandatory=True).model_dump())
    return issues


class ReviewWorkflowService:
    def __init__(self, store: ReviewStore | None = None, provider=None) -> None:
        self.store = store or ReviewStore()
        self.provider = provider

    def _pipeline(self) -> ContentQAPipeline:
        return ContentQAPipeline("data", provider=self.provider)

    def create(self, request: ReviewCreateRequest) -> dict:
        created = now(); session_id = str(uuid.uuid4())
        content = ContentInput(**request.model_dump(exclude={"source_filename"}))
        policy = ReviewPolicyBuilder("data").build(content)
        by_source = {source["source"]: source for source in policy.sources}
        rules = policy.hard_rules + policy.soft_rules
        basis = []
        labels = {"task": "任务要求", "brand": "品牌规则", "product": "产品知识", "platform": "合规规则", "rubric": "内容质量标准", "custom": "自定义规则"}
        for key in ("task", "brand", "product", "platform", "rubric", "custom"):
            source = by_source.get(key, {"version": "1.0"})
            basis.append({"key": key, "label": labels[key], "version": source.get("version", "1.0"), "count": len([r for r in rules if r.source == key])})
        payload = {
            "id": session_id, "state": WorkflowState.NEW_REVIEW, "createdAt": created, "updatedAt": created,
            "sourceFilename": request.source_filename, "taskSummary": f"这是一篇{request.sub_direction}内容，将按照{request.writing_style}风格进行审核。",
            "reviewFocus": ["关键词要求", "品牌名称一致性", "产品事实", "合规表达", "内容准确性", "品牌语气"],
            "reviewBasis": basis, "summary": ReviewSummaryViewModel().model_dump(), "firstReview": [], "optimizedDraft": None,
            "secondReview": [], "finalVersion": None, "error": None, "input": content.model_dump(), "policy": policy.model_dump(),
            "firstDecisions": {}, "secondDecisions": {}, "technical": {},
        }
        self._save(payload); return payload

    def first_review(self, payload: dict) -> dict:
        payload["state"] = WorkflowState.REVIEWING
        result = self._pipeline().run({**payload["input"], "review_policy": payload["policy"]}, review_only=True)
        payload["technical"]["first"] = result.model_dump()
        payload["firstReview"] = self._issues(result, payload["firstDecisions"]) + evidence_issues(payload["input"]["original_text"])
        payload["summary"] = self._summary(result, payload["firstReview"])
        payload["state"] = WorkflowState.WAITING_FOR_FIRST_DECISIONS if payload["firstReview"] else WorkflowState.FIRST_REVIEW_COMPLETE
        self._save(payload); return payload

    def decisions(self, payload: dict, stage: str, decisions: dict[str, str]) -> dict:
        key = "firstDecisions" if stage == "first" else "secondDecisions"
        payload[key].update(decisions)
        result_key = "firstReview" if stage == "first" else "secondReview"
        for issue in payload[result_key]: issue["decision"] = payload[key].get(issue["id"])
        if stage == "first": payload["summary"] = self._summary_from_views(payload, result_key, "firstDecisions")
        self._save(payload); return payload

    def rewrite(self, payload: dict) -> dict:
        self._assert_mandatory(payload, "firstReview", "firstDecisions")
        payload["state"] = WorkflowState.GENERATING_OPTIMIZED_DRAFT
        accepted = [key for key, value in payload["firstDecisions"].items() if value == "accepted"]
        result = self._pipeline().run({**payload["input"], "review_policy": payload["policy"]}, accepted)
        result.final_text = self._conservative_financial_rewrite(payload["input"]["original_text"], result.final_text)
        self._guard_claims(payload["input"]["original_text"], result)
        payload["technical"]["rewrite"] = result.model_dump()
        payload["optimizedDraft"] = self._draft(result)
        payload["state"] = WorkflowState.OPTIMIZED_DRAFT_READY
        self._save(payload); return payload

    def second_review(self, payload: dict) -> dict:
        payload["state"] = WorkflowState.SECOND_REVIEWING
        draft = payload["optimizedDraft"]["content"]
        result = self._pipeline().run({**payload["input"], "original_text": draft, "review_policy": payload["policy"]}, review_only=True)
        payload["technical"]["second"] = result.model_dump()
        payload["secondReview"] = self._issues(result, payload["secondDecisions"])
        payload["state"] = WorkflowState.WAITING_FOR_SECOND_DECISIONS if payload["secondReview"] else WorkflowState.SECOND_REVIEW_COMPLETE
        self._save(payload); return payload

    def finalize(self, payload: dict, keep_current: bool = False) -> dict:
        self._assert_mandatory(payload, "secondReview", "secondDecisions")
        payload["state"] = WorkflowState.GENERATING_FINAL
        if keep_current:
            final = payload["optimizedDraft"]
        else:
            accepted = [key for key, value in payload["secondDecisions"].items() if value == "accepted"]
            result = self._pipeline().run({**payload["input"], "original_text": payload["optimizedDraft"]["content"], "review_policy": payload["policy"]}, accepted)
            result.final_text = self._conservative_financial_rewrite(payload["input"]["original_text"], result.final_text)
            self._guard_claims(payload["input"]["original_text"] + payload["optimizedDraft"]["content"], result)
            payload["technical"]["final"] = result.model_dump(); final = self._draft(result)
        payload["finalVersion"] = final; payload["state"] = WorkflowState.READY_TO_DELIVER
        self._save(payload); return payload

    def export_docx(self, payload: dict) -> bytes:
        from docx import Document
        final = payload.get("finalVersion") or payload.get("optimizedDraft")
        if not final: raise ValueError("Final content is not ready")
        document = Document(); document.add_heading(final["titles"][0] if final["titles"] else "最终版本", 0); document.add_paragraph(final["content"]); document.add_paragraph(" ".join(final["tags"]))
        buffer = io.BytesIO(); document.save(buffer); return buffer.getvalue()

    def _issues(self, result: PipelineResult, decisions: dict[str, str]) -> list[dict]:
        return [issue_view(issue, result.review_policy, decisions).model_dump() for diagnosis in result.diagnoses.values() for issue in diagnosis.issues]

    def _summary(self, result: PipelineResult, issues: list[dict]) -> dict:
        summary = ReviewSummaryViewModel(autoFixedCount=len(result.report.get("hard_rule_fixes", [])), pendingCount=len(issues), complianceStatus="待确认" if issues else "合规检查通过")
        return summary.model_dump()

    def _summary_from_views(self, payload: dict, views_key: str, decisions_key: str) -> dict:
        decisions = payload[decisions_key]; views = payload[views_key]
        return {**payload["summary"], "pendingCount": sum(v["id"] not in decisions for v in views), "acceptedCount": sum(v == "accepted" for v in decisions.values()), "ignoredCount": sum(v == "ignored" for v in decisions.values())}

    def _draft(self, result: PipelineResult) -> dict:
        return {"titles": result.titles, "content": result.final_text, "tags": result.tags, "changes": result.report.get("generation_changes", []), "validation": result.rewrite_validation.model_dump() if result.rewrite_validation else {}}

    def _assert_mandatory(self, payload: dict, views: str, decisions: str) -> None:
        unresolved = [issue["id"] for issue in payload[views] if issue["isMandatory"] and payload[decisions].get(issue["id"]) != "accepted"]
        if unresolved: raise ValueError("Mandatory issues must be accepted or edited before continuing")

    def _guard_claims(self, source: str, result: PipelineResult) -> None:
        claims = unsupported_claims(source, "\n".join(result.titles + [result.final_text]))
        if claims: raise ValueError("Generated content contains unsupported factual claims: " + ", ".join(claims))

    def _conservative_financial_rewrite(self, original: str, candidate: str) -> str:
        if not re.search(r"收益|投资|理财", original):
            return candidate
        candidate = re.sub(r"保证.{0,8}收益率?", "历史收益表现", candidate)
        candidate = re.sub(r"超过?\d+(?:\.\d+)?%的用户[^。！!?]*[。！!?]?", "", candidate)
        if "投资有风险" not in candidate:
            candidate = candidate.rstrip() + "\n\n投资有风险，请审慎决策。"
        return candidate

    def _save(self, payload: dict) -> None:
        payload["updatedAt"] = now(); self.store.save(payload["id"], payload["createdAt"], payload["updatedAt"], payload)
