from __future__ import annotations

import io
from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedDocument:
    filename: str
    text: str
    status: str = "parsed"
    error: str | None = None


class DocumentParser(ABC):
    supported_extensions = {".txt", ".md", ".docx", ".pdf", ".csv", ".xlsx"}

    @abstractmethod
    def parse(self, filename: str, data: bytes) -> ParsedDocument:
        raise NotImplementedError


class DefaultDocumentParser(DocumentParser):
    def parse(self, filename: str, data: bytes) -> ParsedDocument:
        extension = Path(filename).suffix.lower()
        if extension not in self.supported_extensions:
            return ParsedDocument(filename, "", "failed", f"Unsupported file type: {extension}")
        try:
            if extension in {".txt", ".md"}:
                text = data.decode("utf-8-sig")
            elif extension == ".docx":
                from docx import Document
                text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)
            elif extension == ".pdf":
                from pypdf import PdfReader
                text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
            else:
                import pandas as pd
                frame = pd.read_csv(io.BytesIO(data)) if extension == ".csv" else pd.read_excel(io.BytesIO(data))
                text = frame.to_csv(index=False)
            return ParsedDocument(filename, text, "parsed")
        except Exception as error:
            return ParsedDocument(filename, "", "failed", str(error))
